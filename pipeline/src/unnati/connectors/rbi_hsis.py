"""RBI Handbook of Statistics on Indian States: state finances, power, environment and more.

The handbook's index page (rbi.org.in) lists every table with an XLSX link, but the files on
rbidocs.rbi.org.in sit behind a JavaScript bot challenge. The challenge is not bypassed: the
maintainer downloads the needed files in a browser once per edition into
`pipeline/manual-downloads/rbi_hsis/<edition>/` (`unnati manual rbi-links` prints the links), and the
files are committed. RBI names each file after its table ("164T_<id>.XLSX"), so the importer
finds tables by that prefix. The daily run reads the index page to notice a newer edition."""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from itertools import pairwise
from pathlib import Path

from unnati.core.entities import EntityResolver, UnknownEntityError
from unnati.core.http import PoliteClient
from unnati.core.periods import Period, calendar_year, fiscal_year
from unnati.observations import Observation
from unnati.reference import Population

INDEX_URL = "https://www.rbi.org.in/Scripts/AnnualPublications.aspx?head=Handbook%20of%20Statistics%20on%20Indian%20States"
MANUAL_DIR = Path(__file__).resolve().parents[3] / "manual-downloads" / "rbi_hsis"

# Tables the importer reads, by number in the 2024-25 edition (titles are checked on import,
# since numbering can shift between editions).
TABLES = {
    10: "Poverty Estimates - Multi-dimensional Poverty Index",
    16: "Health Infrastructure - Doctors and Specialists",
    17: "Number of Government Hospitals and Beds",
    20: "Per Capita Net State Domestic Product (Constant Prices)",
    21: "Gross State Domestic Product (Current Prices)",
    22: "Gross State Domestic Product (Constant Prices)",
    99: "Forest Cover",
    107: "Status of Ground Water Extraction",
    138: "Per Capita Availability of Power",
    140: "Installed Capacity of Power",
    143: "Installed Capacity of Grid Interactive Renewable Power",
    149: "Telephones per 100 Population",
    153: "Credit-Deposit Ratio of Scheduled Commercial Banks",
    164: "Gross Fiscal Deficit",
    166: "Revenue Expenditure",
    168: "Own Tax Revenue",
    173: "Capital Expenditure",
    174: "Capital Outlay",
    176: "Composition of Outstanding Liabilities",
    181: "State-wise Exports",
}


@dataclass(frozen=True)
class Listed:
    number: int
    title: str
    url: str


def edition_of(page: str) -> str:
    match = re.search(r"Handbook of Statistics on Indian States (\d{4}-\d{2})", page)
    if not match:
        raise RuntimeError("RBI handbook index: edition not found")
    return match.group(1)


def listed_tables(page: str) -> list[Listed]:
    out = []
    for title, rest in re.findall(r"PublicationsView\.aspx\?id=\d+>([^<]+)</a>(.{0,600})", page, re.S):
        number = re.match(r"\s*Table\s*(\d+)\s*:", html.unescape(title))
        links = re.findall(r"href=['\"]?([^'\" >]+\.xlsx)", rest, re.I)
        if number and links:
            out.append(Listed(int(number.group(1)), html.unescape(title).strip(), links[0]))
    return out


def index(http: PoliteClient) -> tuple[str, list[Listed]]:
    page = http.get(INDEX_URL).text
    return edition_of(page), listed_tables(page)


def local_files(edition: str) -> dict[int, Path]:
    """Table number -> downloaded file, from RBI's "<n>T_" file-name prefix."""
    folder = MANUAL_DIR / edition
    found = {}
    for path in folder.glob("*"):
        match = re.match(r"(?i)^(\d+)T_.*\.xlsx$", path.name)
        if match:
            found[int(match.group(1))] = path
    return found


# --- Import ------------------------------------------------------------------------------------

_SKIP_ROWS = re.compile(
    r"(?i)^(northern|north-eastern|eastern|central|western|southern) region|^others$|^all states"
)
_INDIA = re.compile(r"(?i)^(all[- ]india|total)$")
_FY = re.compile(r"^(\d{4})-(\d{2})\s*(?:\((A|RE|BE)\))?$")
_NOTE = re.compile(r"^(?:-|\.|\*|\d+\.\s|Note|Source|SDLs|[A-Z]{1,3}:)")


@dataclass(frozen=True)
class Cell:
    place: str
    column: str
    value: float | None
    flagged: bool  # the source marks the figure with "*" (carried forward, not normalised, etc.)


def _place(raw: object) -> str | None:
    name = re.sub(r"\s*\*+$", "", str(raw or "").strip())
    if not name or _NOTE.match(name) or _SKIP_ROWS.search(name):
        return None
    return "All India" if _INDIA.match(name) else name


def _number(raw: object) -> tuple[float | None, bool]:
    text = str(raw).strip() if raw is not None else ""
    flagged = text.endswith("*")
    text = text.rstrip("*").replace(",", "").strip()
    try:
        return float(text), flagged
    except ValueError:
        return None, flagged


def wide(path: Path, header_pattern: str = r"(?i)^(region/)?state") -> list[Cell]:
    """Cells of a states-by-columns table spread over one or more sheets."""
    import openpyxl

    cells: list[Cell] = []
    for ws in openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets:
        header: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            values = list(row[1:])  # column A is always blank
            first = values[0] if values else None
            if header is None:
                if first is not None and re.match(header_pattern, str(first).strip()):
                    header = [str(c).strip() if c is not None else "" for c in values]
                continue
            if first is None and any(re.fullmatch(r"\d{4}(-\d{2})?.*", str(c or "")) for c in values[1:]):
                # Years on a second header row (e.g. under "Base: 2011-12"): they win. Sub-labels
                # such as "Headcount Ratio" under "NFHS-5" are not years and leave the header alone.
                header = [
                    str(c).strip() if c is not None else h for c, h in zip(values, header, strict=False)
                ]
                continue
            place = _place(first)
            if place is None:
                continue
            for label, raw in zip(header[1:], values[1:], strict=False):
                if label:
                    value, flagged = _number(raw)
                    cells.append(Cell(place, label, value, flagged))
    return cells


def unit_of(path: Path) -> str | None:
    """The money unit printed under a table's title: "lakh" or "crore"."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    for row in ws.iter_rows(max_row=6, values_only=True):
        for cell in row:
            match = re.search(r"₹\s*(lakh|crore)", str(cell or ""), re.I)
            if match:
                return match.group(1).lower()
    return None


def fiscal_column(label: str) -> tuple[Period, str] | None:
    """("2023-24 (RE)") -> (FY 2023-24, "RE"); "A" when unmarked (older years are actuals)."""
    match = _FY.match(label.strip())
    if not match:
        return None
    return fiscal_year(int(match.group(1))), match.group(3) or "A"


def liabilities(path: Path) -> list[tuple[str, Period, float]]:
    """(place, FY, total outstanding liabilities) from Table 176: one sheet pair per end-March
    year, with the total in the "Outstanding Liabilities" column of the second sheet."""
    import openpyxl

    out = []
    for ws in openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets:
        rows = list(ws.iter_rows(values_only=True))
        title = next((str(r[1]) for r in rows if len(r) > 1 and r[1] and "TABLE 176" in str(r[1])), "")
        year = re.search(r"- (\d{4})", title)
        header_at = next(
            (i for i, r in enumerate(rows) if len(r) > 1 and r[1] and str(r[1]).startswith("State")), None
        )
        if not year or header_at is None:
            continue
        header = [str(c).strip() if c else "" for c in rows[header_at]]
        if "Outstanding Liabilities" not in header:
            continue
        column = header.index("Outstanding Liabilities")
        fy = fiscal_year(int(year.group(1)) - 1)  # end-March 2025 closes FY 2024-25
        for row in rows[header_at + 1 :]:
            place = _place(row[1])
            value, _ = _number(row[column])
            if place and value is not None:
                out.append((place, fy, value))
    return out


class _Builder:
    def __init__(self, edition: str, resolver: EntityResolver) -> None:
        self.source = f"RBI Handbook of Statistics on Indian States {edition}"
        self.resolver = resolver
        self.out: list[Observation] = []
        self.problems: list[str] = []

    def slug(self, place: str, period: Period) -> str | None:
        try:
            entity = self.resolver.resolve(place, period)
        except UnknownEntityError as err:
            if "none is valid for" in str(err):
                # RBI keeps rows for places outside their lifetime (e.g. Daman & Diu after the 2020
                # merger, usually blank or repeated); those cells are skipped.
                self.problems.append(f"{self.source}: skipped {place!r} outside its validity")
            else:
                self.problems.append(f"{self.source}: {err}")
            return None
        return entity.slug if entity else None

    def fiscal(self, path: Path, statuses: tuple[str, ...] = ("A", "RE")) -> dict[tuple[str, str], tuple]:
        """(slug, FY label) -> (value, period, status) for a table with fiscal-year columns."""
        found = {}
        for cell in wide(path):
            parsed = fiscal_column(cell.column)
            if parsed is None or cell.value is None or parsed[1] not in statuses:
                continue
            period, status = parsed
            slug = self.slug(cell.place, period)
            if slug:
                found[(slug, period.label)] = (cell.value, period, status)
        return found

    def add(self, indicator_id: str, slug: str, period: Period, value: float, note: str, provisional=False):
        self.out.append(
            Observation(
                indicator_id,
                slug,
                period,
                round(value, 3),
                is_provisional=provisional,
                note=f"{self.source}, {note}",
            )
        )


def _economy_fill_in(b: _Builder, files: dict[int, Path], mospi_covered: set[str]) -> None:
    """Per-capita NSDP (Table 20) and real GSDP growth from GSDP at constant prices (Table 22),
    both base 2011-12 like MoSPI's series, for places missing from MoSPI's API."""
    from unnati.connectors.mospi import three_year_growth

    if 20 in files:
        for (slug, _), (value, period, _) in b.fiscal(files[20], ("A",)).items():
            if slug not in mospi_covered and period.start.year >= 2011:
                b.add(
                    "per-capita-nsdp-constant",
                    slug,
                    period,
                    value,
                    "table 20 (per capita NSDP, base 2011-12)",
                )
    # GSDP totals (Tables 21 and 22 are in ₹ lakh; the indicators are in ₹ crore like MoSPI's).
    for table, indicator_id, what in ((21, "gsdp-current", "current"), (22, "gsdp-constant", "constant")):
        if table in files:
            to_crore = 1 / 100 if unit_of(files[table]) == "lakh" else 1.0
            for (slug, _), (value, period, _) in b.fiscal(files[table], ("A",)).items():
                if slug not in mospi_covered and period.start.year >= 2011:
                    note = f"table {table} (GSDP at {what} prices, base 2011-12)"
                    b.add(indicator_id, slug, period, value * to_crore, note)
    if 22 in files:
        series: dict[str, list[tuple[Period, float]]] = {}
        for (slug, _), (value, period, _) in b.fiscal(files[22], ("A",)).items():
            if slug not in mospi_covered:
                series.setdefault(slug, []).append((period, value))
        growth: list[Observation] = []
        for slug, points in series.items():
            points.sort(key=lambda p: p[0].start)
            for (p0, v0), (p1, v1) in pairwise(points):
                if p1.start.year == p0.start.year + 1 and v0 > 0 and p1.start.year >= 2012:
                    note = f"{b.source}, table 22: growth of GSDP at constant prices (base 2011-12)"
                    growth.append(
                        Observation("gsdp-growth-real", slug, p1, round((v1 / v0 - 1) * 100, 2), note=note)
                    )
        b.out.extend(growth)
        b.out.extend(three_year_growth(growth))


def observations(
    files: dict[int, Path],
    edition: str,
    resolver: EntityResolver,
    population: dict[tuple[str, int], Population],
    gsdp_fallback: dict[tuple[str, str], float],
    mospi_covered: set[str] = frozenset(),
) -> tuple[list[Observation], list[str]]:
    """Indicators from the downloaded tables. `gsdp_fallback` maps (entity slug, FY label) to GSDP
    at current prices in ₹ crore (MoSPI), used where Table 21 has no value. Per-capita income and
    real growth (Tables 20 and 22) are only added for places MoSPI's API doesn't cover
    (`mospi_covered`), so the two sources never disagree on the same series."""
    b = _Builder(edition, resolver)
    _economy_fill_in(b, files, mospi_covered)
    missing = [n for n in TABLES if n not in files]
    if missing:
        b.problems.append(f"{b.source}: tables not downloaded: {', '.join(map(str, missing))}")

    # GSDP denominators: RBI Table 21 when present (it covers every state), else MoSPI.
    gsdp: dict[tuple[str, str], float] = {}
    if 21 in files:
        to_crore = 1 / 100 if unit_of(files[21]) == "lakh" else 1.0  # Table 21 is in ₹ lakh
        gsdp = {key: value * to_crore for key, (value, *_) in b.fiscal(files[21], ("A", "RE")).items()}
    for key, value in gsdp_fallback.items():
        gsdp.setdefault(key, value)

    # State finances from FY 2014-15 (current boundaries): actuals, and revised estimates as
    # provisional; budget estimates are skipped.
    for table, indicator_id, what in (
        (164, "fiscal-deficit-pct-gsdp", "gross fiscal deficit"),
        (168, "own-tax-revenue-pct-gsdp", "own tax revenue"),
    ):
        if table in files:
            for (slug, fy), (value, period, status) in b.fiscal(files[table]).items():
                denominator = gsdp.get((slug, fy))
                if period.start.year >= 2014 and denominator:
                    note = f"table {table} ({what}, {status}) / GSDP at current prices"
                    b.add(indicator_id, slug, period, value / denominator * 100, note, status == "RE")
    if 166 in files and 174 in files:
        revenue, outlay = b.fiscal(files[166]), b.fiscal(files[174])
        for key, (value, period, status) in outlay.items():
            if key in revenue and period.start.year >= 2014:
                share = value / (revenue[key][0] + value) * 100
                note = f"tables 174/166: capital outlay / (revenue expenditure + capital outlay), {status}"
                b.add("capex-share", key[0], period, share, note, "RE" in (status, revenue[key][2]))
    if 176 in files:
        rows = liabilities(files[176])
        series = {(place, fy.start.year): value for place, fy, value in rows}
        for place, fy, value in rows:
            before, after = series.get((place, fy.start.year - 1)), series.get((place, fy.start.year + 1))
            if before and after and value > 2.5 * max(before, after):
                # e.g. Tripura end-March 2015: total internal debt is ten times the sum of its parts
                b.problems.append(
                    f"{b.source}: table 176 {place} {fy.label} skipped as a likely typo ({value:,.0f})"
                )
                continue
            slug = b.slug(place, fy)
            denominator = gsdp.get((slug, fy.label)) if slug else None
            if slug and fy.start.year >= 2014 and denominator:
                note = "table 176 (outstanding liabilities, end-March) / GSDP at current prices"
                b.add("debt-pct-gsdp", slug, fy, value / denominator * 100, note)

    if 138 in files:
        for (slug, _), (value, period, _) in b.fiscal(files[138], ("A",)).items():
            b.add("electricity-per-capita", slug, period, value, "table 138 (per capita availability, CEA)")

    if 140 in files and 143 in files:  # renewables at end-March / installed capacity in the FY to that March
        installed = b.fiscal(files[140], ("A",))
        for cell in wide(files[143]):
            if not re.fullmatch(r"\d{4}", cell.column) or cell.value is None:
                continue
            fy = fiscal_year(int(cell.column) - 1)
            slug = b.slug(cell.place, fy)
            total = installed.get((slug, fy.label)) if slug else None
            if total and total[0] > 0:
                note = "tables 143/140: grid-interactive renewables (excl. large hydro) / installed capacity"
                b.add("renewable-share-capacity", slug, fy, min(cell.value / total[0] * 100, 100), note)

    if 107 in files:
        for cell in wide(files[107]):
            year = re.search(r"Stage of Ground Water Extraction.*\((\d{4})\)", cell.column)
            if not year or cell.value is None or cell.flagged:  # "*": carried over from an older round
                continue
            period = calendar_year(int(year.group(1)))
            slug = b.slug(cell.place, period)
            if slug:
                b.add("groundwater-extraction-stage", slug, period, cell.value, "table 107 (CGWB assessment)")

    if 99 in files:  # % change in forest cover between comparable ISFR rounds
        rounds: dict[str, list[tuple[int, float, bool]]] = {}
        for cell in wide(files[99]):
            year = re.fullmatch(r"(\d{4})(\*?)", cell.column.strip())
            if year and cell.value:
                rounds.setdefault(cell.place, []).append(
                    (int(year.group(1)), cell.value, bool(year.group(2)))
                )
        for place, series in rounds.items():
            series.sort()
            for (y0, a, star0), (y1, value, star1) in pairwise(series):
                # FSI's 2013, 2021 and 2023 figures are not normalised like the rest; older rounds
                # used coarser imagery, so only changes from 2011 on are kept.
                if star0 != star1 or y0 < 2011:
                    continue
                period = calendar_year(y1)
                slug = b.slug(place, period)
                # J&K's series mixes the UT with and without Ladakh between rounds.
                if slug and slug not in ("jammu-and-kashmir", "ladakh", "jammu-and-kashmir-state"):
                    note = f"table 99: forest cover in ISFR {y1} vs ISFR {y0}"
                    b.add("forest-cover-change", slug, period, (value - a) / a * 100, note)

    if 181 in files:  # merchandise exports by state of origin, US$ million per FY
        for cell in wide(files[181]):
            parsed = fiscal_column(cell.column)
            if parsed is None or cell.value is None:
                continue
            period = parsed[0]
            slug = b.slug(cell.place, period)
            people = population.get((slug, period.start.year)) if slug else None
            if people:
                note = "table 181 (DGCI&S exports by state of origin) / projected population"
                b.add("exports-per-capita", slug, period, cell.value * 1e6 / people.persons, note)

    if 10 in files:  # MPI headcount ratio, NFHS-4 and NFHS-5 rounds
        from unnati.connectors.mospi import NFHS_ROUNDS

        for cell in wide(files[10]):
            round_ = re.match(r"(NFHS-\d)", cell.column)
            period = NFHS_ROUNDS.get(round_.group(1).lower()) if round_ else None
            if period and cell.value is not None:
                slug = b.slug(cell.place, period)
                if slug:
                    b.add("mpi-headcount", slug, period, cell.value, "table 10 (NITI Aayog National MPI)")
    return b.out, sorted(set(b.problems))
